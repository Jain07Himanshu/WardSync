"""Tests for document parsers (PDF, DOCX, CSV, HTML, TXT)."""

import io
import docx
from pypdf import PdfWriter
from app.ingestion.parsers import (
    parse_plain_text,
    parse_csv,
    parse_html_content,
    parse_docx,
    parse_pdf,
)


def test_parse_plain_text():
    raw = b"Reported water leakage in Andheri West."
    parsed = parse_plain_text(raw)
    assert parsed == "Reported water leakage in Andheri West."


def test_parse_csv():
    csv_bytes = b"ComplaintID,Category,Location\n101,Roads,S.V. Road\n102,Water,Milan Subway"
    parsed = parse_csv(csv_bytes)
    assert "ComplaintID" in parsed
    assert "S.V. Road" in parsed
    assert "Milan Subway" in parsed


def test_parse_html_content():
    html_sample = """
    <!DOCTYPE html>
    <html>
    <head><title>Civic Dispatch: Pothole Repairs Underway</title></head>
    <body>
        <nav><a href="/">Home</a></nav>
        <article>
            <h1>Andheri West Road Work</h1>
            <p>Municipal crew deployed asphalt resurfacing on S.V. Road after citizen grievances.</p>
        </article>
        <footer>Copyright 2026</footer>
    </body>
    </html>
    """
    title, body = parse_html_content(html_sample)
    assert "Civic Dispatch" in title
    assert "Municipal crew deployed asphalt resurfacing" in body
    assert "Home" not in body  # Navigation stripped
    assert "Copyright 2026" not in body  # Footer stripped


def test_parse_docx():
    doc = docx.Document()
    doc.add_paragraph("Municipal Drainage Inspection Report")
    doc.add_paragraph("Storm water drain cleared near station approach.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    parsed = parse_docx(docx_bytes)
    assert "Municipal Drainage Inspection Report" in parsed
    assert "Storm water drain cleared" in parsed


def test_parse_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    parsed = parse_pdf(pdf_bytes)
    # Blank page returns either empty or note
    assert isinstance(parsed, str)
